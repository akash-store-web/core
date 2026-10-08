"""Lee la hoja del catalogo real de Akash Store (T-01 #1) y muestra que se cargaria.

PASO 1: este script NO toca la base de datos. Solo lee el Excel, revisa que cada fila cumpla
las reglas del backend (ver docs/API.md) y escribe un resumen. Si algo esta mal, lo dice.

Uso (desde la carpeta backend/, con el entorno virtual activo):
    python -m scripts.cargar_catalogo --archivo "D:\\ruta\\catalogo_akashstore_2025.xlsx"

Reglas que revisa:
- La categoria debe ser una de las 7 oficiales.
- El producto necesita descripcion y precio base mayor que 0 (maximo 2 decimales).
  Sin precio base NO se puede crear: se salta y se avisa.
- Nombre del producto: maximo 120 caracteres. Material, medidas y variante: maximo 80.
- Todas las filas de un mismo producto deben tener la misma descripcion, material, medidas,
  precio base y valor de "es_pieza_natural".
- No puede haber dos variantes con el mismo nombre dentro de un producto.
- El precio de una variante, si existe, debe ser mayor que 0.
- Columna opcional "existencias": numero entero mayor o igual a 0. Si no existe, todo entraria con 0.
"""
import argparse
import sys
import unicodedata
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

CATEGORIAS = {
    "Boxes y kits", "Pulseras", "Collares y colgantes", "Anillos",
    "Roll-on", "Inciensos y limpieza energética", "Decoración y amuletos",
}
COLUMNAS_OBLIGATORIAS = ["categoria", "producto", "descripcion", "precio_base", "es_pieza_natural", "variante"]
MAX_NOMBRE = 120
MAX_CORTO = 80


def texto(valor) -> str:
    return "" if valor is None else str(valor).strip()


def sin_tildes(cadena: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", cadena) if unicodedata.category(c) != "Mn").lower()


def a_decimal(valor, campo: str, contexto: str, errores: list):
    """Devuelve un Decimal > 0 con maximo 2 decimales, o None si la celda esta vacia."""
    s = texto(valor)
    if not s:
        return None
    try:
        d = Decimal(s.replace(",", "."))
    except InvalidOperation:
        errores.append(f"{contexto}: {campo} no es un numero ({s!r})")
        return None
    if d <= 0:
        errores.append(f"{contexto}: {campo} debe ser mayor que 0 ({s})")
        return None
    if d.as_tuple().exponent < -2:
        errores.append(f"{contexto}: {campo} tiene mas de 2 decimales ({s})")
        return None
    return d


def leer_hoja(ruta: str):
    libro = load_workbook(ruta, data_only=True, read_only=True)
    hoja = libro["Catalogo_2025"] if "Catalogo_2025" in libro.sheetnames else libro.worksheets[0]
    filas = hoja.iter_rows(values_only=True)
    encabezados = [sin_tildes(texto(c)) for c in next(filas)]
    faltan = [c for c in COLUMNAS_OBLIGATORIAS if c not in encabezados]
    if faltan:
        raise SystemExit(f"Faltan columnas en la hoja: {', '.join(faltan)}")
    registros = []
    for numero, fila in enumerate(filas, start=2):
        reg = dict(zip(encabezados, fila))
        if not texto(reg.get("producto")):
            continue
        reg["_fila"] = numero
        registros.append(reg)
    return encabezados, registros


def agrupar(registros):
    productos = {}
    for reg in registros:
        productos.setdefault(texto(reg["producto"]), []).append(reg)
    return productos


def revisar(nombre: str, filas: list) -> dict:
    errores, notas_confirmar = [], []
    ctx = f"{nombre!r}"
    primera = filas[0]

    if len(nombre) > MAX_NOMBRE:
        errores.append(f"{ctx}: el nombre pasa de {MAX_NOMBRE} caracteres")

    categoria = texto(primera.get("categoria"))
    if categoria not in CATEGORIAS:
        errores.append(f"{ctx}: categoria no oficial ({categoria!r})")

    for campo in ("descripcion", "material", "medidas", "precio_base", "es_pieza_natural"):
        if campo in primera and len({texto(f.get(campo)) for f in filas}) > 1:
            errores.append(f"{ctx}: la columna {campo} cambia entre las filas del mismo producto")

    if not texto(primera.get("descripcion")):
        errores.append(f"{ctx}: falta la descripcion")
    for campo in ("material", "medidas"):
        if len(texto(primera.get(campo))) > MAX_CORTO:
            errores.append(f"{ctx}: {campo} pasa de {MAX_CORTO} caracteres")

    natural = sin_tildes(texto(primera.get("es_pieza_natural")))
    if natural not in ("si", "no"):
        errores.append(f"{ctx}: es_pieza_natural debe decir Si o No ({texto(primera.get('es_pieza_natural'))!r})")

    precio_base = a_decimal(primera.get("precio_base"), "precio_base", ctx, errores)

    variantes, vistos = [], set()
    for f in filas:
        nombre_var = texto(f.get("variante")) or "Única"
        c = f"{ctx} / variante {nombre_var!r}"
        if len(nombre_var) > MAX_CORTO:
            errores.append(f"{c}: el nombre pasa de {MAX_CORTO} caracteres")
        if nombre_var.lower() in vistos:
            errores.append(f"{c}: variante repetida")
        vistos.add(nombre_var.lower())
        precio_var = a_decimal(f.get("precio_variante"), "precio_variante", c, errores)
        existencias = 0
        if "existencias" in f and texto(f.get("existencias")):
            try:
                existencias = int(Decimal(texto(f["existencias"])))
                if existencias < 0:
                    raise ValueError
            except (InvalidOperation, ValueError):
                errores.append(f"{c}: existencias debe ser un numero entero >= 0")
                existencias = 0
        variantes.append({"nombre": nombre_var, "precio": precio_var, "existencias": existencias})
        nota = texto(f.get("notas"))
        if "confirm" in nota.lower():
            notas_confirmar.append(nota)

    if errores:
        estado = "ERROR"
    elif precio_base is None:
        estado = "SALTA"
    else:
        estado = "OK"
    return {
        "nombre": nombre, "categoria": categoria, "precio_base": precio_base,
        "variantes": variantes, "estado": estado, "errores": errores,
        "confirmar": list(dict.fromkeys(notas_confirmar)),
    }


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    parser = argparse.ArgumentParser(description="Revisa la hoja del catalogo real (no escribe en la base).")
    parser.add_argument("--archivo", required=True, help="ruta del Excel del catalogo")
    args = parser.parse_args()

    encabezados, registros = leer_hoja(args.archivo)
    productos = agrupar(registros)
    resultados = [revisar(nombre, filas) for nombre, filas in productos.items()]

    print(f"Hoja leida: {len(registros)} filas, {len(productos)} productos.\n")
    categoria_actual = None
    for r in resultados:
        if r["categoria"] != categoria_actual:
            categoria_actual = r["categoria"]
            print(f"== {categoria_actual}")
        precio = f"S/ {r['precio_base']:.2f}" if r["precio_base"] is not None else "SIN PRECIO"
        print(f"  [{r['estado']}] {r['nombre']} - {precio} - {len(r['variantes'])} variante(s)")

    ok = [r for r in resultados if r["estado"] == "OK"]
    saltados = [r for r in resultados if r["estado"] == "SALTA"]
    con_error = [r for r in resultados if r["estado"] == "ERROR"]

    print("\n---- RESUMEN ----")
    print(f"Se cargarian: {len(ok)} productos y {sum(len(r['variantes']) for r in ok)} variantes.")
    if saltados:
        print(f"Se saltan por no tener precio base ({len(saltados)}):")
        for r in saltados:
            print(f"  - {r['nombre']}")
    if con_error:
        print(f"Con errores ({len(con_error)}):")
        for r in con_error:
            for e in r["errores"]:
                print(f"  - {e}")
    por_confirmar = [r for r in resultados if r["confirmar"]]
    if por_confirmar:
        print(f"Con notas de 'confirmar' en la hoja ({len(por_confirmar)}): revisarlas con la propietaria antes de cargar.")
    if "existencias" not in encabezados:
        print("AVISO: la hoja no tiene columna 'existencias'. Todo entraria con 0 y se veria 'Agotado'.")
    print("\nNo se escribio nada en la base de datos.")
    sys.exit(1 if con_error else 0)


if __name__ == "__main__":
    main()
